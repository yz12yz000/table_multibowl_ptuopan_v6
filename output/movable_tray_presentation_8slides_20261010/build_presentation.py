"""Build an editable English methods deck with plain-language Chinese notes."""
import json
import math
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
ASSETS = OUT / 'assets'
ASSETS.mkdir(parents=True, exist_ok=True)
OLD = ROOT / 'output/bowl_research_presentation_20261002/assets'
EVIDENCE = ROOT / 'debug_output/clear_table_validation/yolo_20261010'
SUMMARY = json.loads((EVIDENCE / 'summary.json').read_text())
DEMO = json.loads((EVIDENCE / 'offline_demo/result.json').read_text())
for name in ('raw_table.png', 'detection.png', 'grasp_pose.png'):
    shutil.copy2(OLD / name, ASSETS / name)

BG, NAVY, INK, GRAY = 'F4F7FA', '112B3B', '203949', '627785'
TEAL, BLUE, AMBER, RED = '008F83', '3579C7', 'D69826', 'BF5A64'
WHITE, LINE, PALE, TINT = 'FFFFFF', 'D9E3E9', 'E6F3F0', 'EAF1F9'
FONT = 'Lato'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
prs.core_properties.title = 'Movable-Tray Table Clearing — Methods and Validation'
prs.core_properties.subject = 'English slides with accessible Chinese speaker notes'
prs.core_properties.author = 'Robot Grasp Project'
notes = []

def rgb(c):
    return RGBColor.from_string(c)

def shape(s, x, y, w, h, fill=WHITE, stroke=None, kind=MSO_SHAPE.RECTANGLE, lw=1):
    obj = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    obj.fill.solid(); obj.fill.fore_color.rgb = rgb(fill)
    if stroke:
        obj.line.color.rgb = rgb(stroke); obj.line.width = Pt(lw)
    else:
        obj.line.fill.background()
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        obj.adjustments[0] = .12
    style = obj._element.find(qn('p:style'))
    if style is not None:
        effect = style.find(qn('a:effectRef'))
        if effect is not None: effect.set('idx', '0')
    return obj

def text(s, value, x, y, w, h, size=20, color=INK, bold=False, align=None, font=FONT):
    obj = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = obj.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.015)
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.font.name = font; p.font.size = Pt(size)
        p.font.bold = bold; p.font.color.rgb = rgb(color)
        p.space_after = Pt(5); p.line_spacing = 1.08
        if align is not None: p.alignment = align
    return obj

def line(s, x1, y1, x2, y2, color=LINE, width=1.5, arrow=False, dash=False):
    obj = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    obj.line.color.rgb = rgb(color); obj.line.width = Pt(width)
    if arrow:
        tail = OxmlElement('a:tailEnd'); tail.set('type', 'triangle')
        obj._element.spPr.get_or_add_ln().append(tail)
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        obj.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return obj

def tag(s, label, x, y, w, color=TEAL, fill=PALE):
    shape(s, x, y, w, .32, fill)
    text(s, label, x+.08, y+.055, w-.16, .23, 10, color, True)

def bullet(s, title, body, x, y, w=4.6, color=TEAL):
    shape(s, x, y+.06, .055, .28, color)
    text(s, title, x+.18, y, w-.18, .36, 20, INK, True)
    text(s, body, x+.18, y+.47, w-.18, .77, 16, GRAY)

def photo(s, filename, x, y, w, h):
    from PIL import Image
    path = ASSETS / filename
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    return s.shapes.add_picture(str(path), Inches(x+(w-iw*scale)/2), Inches(y+(h-ih*scale)/2),
                               width=Inches(iw*scale), height=Inches(ih*scale))

def note(s, title, body, sources):
    value = '【中文讲解备注｜面向非专业听众】\n' + body + '\n\n【资料来源】\n' + '\n'.join(sources)
    tf = s.notes_slide.notes_text_frame; tf.text = value
    for p in tf.paragraphs:
        p.font.name = 'Noto Sans CJK SC'; p.font.size = Pt(12)
    notes.append((title, value))

def new(title, section, caption='', dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(NAVY if dark else BG)
    text(s, section.upper(), .57, .3, 11.8, .23, 10, '60CEC3' if dark else TEAL, True)
    text(s, title, .57, .83, 12.15, .83, 31, WHITE if dark else INK, True)
    line(s, .57, 7.07, 12.77, 7.07, '36515F' if dark else LINE, .8)
    text(s, caption or 'Movable-Tray Table Clearing  |  Methods & principles', .57, 7.17, 11.3, .19,
         8.5, 'ADBFCA' if dark else GRAY)
    text(s, f'{len(prs.slides):02d}', 12.2, 7.13, .57, .24, 11,
         'ADBFCA' if dark else GRAY, align=PP_ALIGN.RIGHT)
    return s

def disk(s, x, y, r=.25, color=BLUE, label='', fill=None):
    shape(s, x-r, y-r, 2*r, 2*r, fill or WHITE, color, MSO_SHAPE.OVAL, 2)
    if label: text(s, label, x-r, y-.12, 2*r, .25, 12, color, True, PP_ALIGN.CENTER)

def tabletop(s, x, y, w, h, occupied=1, outside=3, selected=None):
    shape(s, x, y, w, h, WHITE, LINE, MSO_SHAPE.ROUNDED_RECTANGLE)
    tx, ty, tw, th = x+w*.49, y+h*.23, w*.43, h*.64
    shape(s, tx, ty, tw, th, PALE, TEAL, MSO_SHAPE.ROUNDED_RECTANGLE, 1.5)
    text(s, 'TRAY', tx+.13, ty+.08, tw-.26, .25, 10, TEAL, True)
    locs = [(tx+tw*.28,ty+th*.39),(tx+tw*.72,ty+th*.39),
            (tx+tw*.28,ty+th*.76),(tx+tw*.72,ty+th*.76)]
    for j,(cx,cy) in enumerate(locs):
        color = TEAL if j < occupied else (AMBER if j == selected else LINE)
        disk(s,cx,cy,min(w*.044,.23),color,chr(65+j),PALE if j<occupied else WHITE)
    for j in range(outside):
        disk(s,x+w*(.16+.12*(j%2)),y+h*(.31+.23*j),min(w*.045,.24),BLUE,str(j+1))
    return locs

def make_demo_charts():
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.patches import Polygon, Circle
    import numpy as np
    plt.rcParams['font.family'] = 'DejaVu Sans'
    for idx in range(1,6):
        observation = json.loads((EVIDENCE / f'offline_demo/round_{idx:04d}/synthetic_observation.json').read_text())
        d = DEMO['rounds'][idx-1]['decision']
        t = observation['tray']; c=np.array(t['center_xy']); theta=t['yaw']
        rot=np.array([[math.cos(theta),-math.sin(theta)],[math.sin(theta),math.cos(theta)]])
        poly=np.array([[-.3,-.22],[.3,-.22],[.3,.22],[-.3,.22]])@rot.T+c
        fig,ax=plt.subplots(figsize=(3.2,3.2),dpi=180)
        fig.patch.set_facecolor('#FFFFFF'); ax.set_facecolor('#FFFFFF')
        ax.add_patch(Polygon(poly,facecolor='#E6F3F0',edgecolor='#008F83',lw=2))
        for obj in observation['objects']:
            p=np.array(obj['center_xy'])
            uv=(p-c)@rot
            inside=abs(uv[0])+.06 <= .3 and abs(uv[1])+.06 <= .22
            ax.add_patch(Circle(p,.06,facecolor='#008F83' if inside else '#3579C7',alpha=.9))
        if d.get('slot'):
            ax.add_patch(Circle(d['slot']['xy'],.073,facecolor='none',edgecolor='#D69826',lw=2.5,ls='--'))
            ax.text(*d['slot']['xy'],d['slot']['id'],ha='center',va='center',fontsize=10,color='#99620B',fontweight='bold')
        ax.set_xlim(-.98,.06); ax.set_ylim(-.83,.50); ax.set_aspect('equal')
        ax.axis('off'); fig.tight_layout(pad=.05)
        for ext in ('png','svg'):
            fig.savefig(ASSETS/f'observation_{idx}.{ext}',bbox_inches='tight',pad_inches=.025)
        plt.close(fig)

make_demo_charts()

s = new('Movable-tray clearing: one bowl at a time', '01 / Goal & feedback loop')
items=[('01','Observe','Acquire a fresh RGB-D frame.'),('02','Measure','Locate the tray and all bowls.'),
       ('03','Verify','Check the previous placement.'),('04','Choose','One source bowl + one free slot.'),
       ('05','Pick & place','Run the existing motion sequence.'),('06','Return','Go back to the saved home pose.')]
coords=[(.72,2.0),(4.86,2.0),(9.,2.0),(9.,4.36),(4.86,4.36),(.72,4.36)]
for (n,t,b),(x,y) in zip(items,coords):
    shape(s,x,y,3.59,1.7,WHITE,LINE)
    text(s,n,x+.2,y+.2,.5,.3,13,TEAL,True)
    text(s,t,x+.2,y+.66,3.18,.4,22,INK,True)
    text(s,b,x+.2,y+1.15,3.18,.43,14,GRAY)
for a,b in [((4.36,2.83),(4.75,2.83)),((8.5,2.83),(8.89,2.83)),
            ((10.79,3.78),(10.79,4.22)),((8.89,5.18),(8.5,5.18)),
            ((4.75,5.18),(4.36,5.18)),((2.5,4.22),(2.5,3.78))]:
    line(s,*a,*b,TEAL,2.3,True)
shape(s,.73,6.32,11.9,.48,PALE)
text(s,'Fresh scene each round  •  Same grasp and motion  •  Move the tray only between actions',
     .88,6.44,11.6,.29,16,TEAL,True)
note(s,'Movable-tray clearing: one bowl at a time',
     '这张图是整个程序的工作循环。先用新图像看桌面，再计算托盘和碗的位置；如果上一轮刚放过一只碗，就先确认它真的进了托盘。确认后，再选下一只和一个空位置，按原来的方法抓放，最后返回原起点。\n“新图像”很重要：上一轮动作结束后，桌面已经变了，旧图像里的空位可能已经放了碗，旧目标的位置也可能不适用了。第一轮没有上一只碗，不需要这一步归位核查。最后一只放完也不能直接结束，还要重新看；默认还会做连续的空源区确认。核查与下一轮选择共用同一次观察，不额外设计一套机械臂检查动作。',
     ['tools/test_live_place_pipeline.py: run_clear_table','tools/place_rounds.py: ClearTableSession.observe'])

s = new('YOLO locates the tray; depth gives its physical position', '02 / Perception & tray localization',
        'Dedicated tray segmentation weights required • outline and coordinate axes are schematic')
shape(s,.72,1.97,6.2,4.7,WHITE,LINE)
text(s,'YOLO OUTER CONTOUR',1.02,2.2,5.65,.27,12,BLUE,True)
shape(s,1.6,2.75,4.7,2.75,TINT,BLUE,MSO_SHAPE.ROUNDED_RECTANGLE,2)
shape(s,1.87,3.02,4.16,2.21,PALE,TEAL,MSO_SHAPE.ROUNDED_RECTANGLE,1.5)
text(s,'USABLE INTERIOR',2.05,3.19,2.2,.25,11,TEAL,True)
for xx,yy in [(1.6,2.75),(6.3,2.75),(6.3,5.5),(1.6,5.5)]:
    disk(s,xx,yy,.055,BLUE,fill=BLUE)
disk(s,4.11,3.92,.08,TEAL,fill=TEAL)
line(s,4.11,3.92,5.35,3.92,TEAL,2.4,True)
line(s,4.11,3.92,4.11,3.28,TEAL,2.4,True)
text(s,'u',5.43,3.8,.35,.35,19,TEAL,True);text(s,'v',4.3,3.17,.35,.35,19,TEAL,True)
disk(s,5.25,4.72,.17,AMBER,'A')
text(s,'Center',3.39,4.18,1.22,.3,13,GRAY)
text(s,'Bowls: YOLO mask + depth → footprint',1.0,5.92,5.66,.37,17,INK,True)
text(s,'Tray: YOLO mask + depth → center and heading',1.0,6.37,5.66,.27,14,GRAY)
bullet(s,'Segment the tray directly','A separate YOLO model supplies the tray outline.',7.5,2.03,5.12)
bullet(s,'Recover XY and heading','Same-frame depth + camera-to-base transform.',7.5,3.46,5.12)
bullet(s,'Move the slots with the tray','p_slot = c_tray + R(θ) · [u, v]',7.5,4.9,5.12)
note(s,'YOLO locates the tray; depth gives its physical position', '',
     ['src/yolo_vision/yolo_vision/tray_utils.py: detect_tray, slot_xy','src/yolo_vision/config/clear_table_tray.json'])


s = new('Pick outside the tray. Place only in verified free space.', '03 / Target & slot selection')
shape(s,.72,1.94,5.86,4.85,WHITE,LINE)
shape(s,6.81,1.94,5.81,4.85,WHITE,LINE)
text(s,'1  SELECT THE BOWL',.98,2.16,5.18,.3,13,BLUE,True)
text(s,'2  SELECT THE SLOT',7.07,2.16,5.12,.3,13,TEAL,True)
shape(s,3.07,2.79,2.93,1.79,PALE,TEAL,MSO_SHAPE.ROUNDED_RECTANGLE,1.7)
text(s,'TRAY',3.24,2.94,1.0,.22,10,TEAL,True)
disk(s,1.69,3.64,.36,BLUE,'OUT',TINT)
disk(s,4.64,3.64,.36,TEAL,'IN',PALE)
disk(s,3.08,4.13,.35,AMBER,'?','FBF1DD')
text(s,'Outside + in source region → pick candidate',.99,4.94,5.22,.38,17,BLUE,True)
text(s,'Fully inside → occupancy only; never re-pick',.99,5.51,5.22,.39,17,TEAL,True)
text(s,'Boundary / unclear depth → observe again',.99,6.09,5.22,.39,16,GRAY)
shape(s,7.2,2.79,5.02,1.79,PALE,TEAL,MSO_SHAPE.ROUNDED_RECTANGLE,1.7)
for xx,label,col,fill,caption in [(7.88,'A',TEAL,PALE,'occupied'),(9.1,'B',AMBER,'FBF1DD','selected'),(10.32,'C',GRAY,'E9EEF2','reserved'),(11.54,'D',BLUE,WHITE,'free')]:
    disk(s,xx,3.48,.28,col,label,fill)
    text(s,caption,xx-.52,4.02,1.04,.24,12,col,align=PP_ALIGN.CENTER)
text(s,'Edge clearance ≥ bowl radius + margin',7.08,4.94,5.24,.38,17,INK,True)
text(s,'Bowl spacing > both radii + clearance',7.08,5.51,5.24,.39,17,INK,True)
text(s,'Keep used slots reserved if detections drop out.',7.08,6.09,5.24,.39,16,GRAY)
note(s,'Pick outside the tray. Place only in verified free space.', '',
     ['src/yolo_vision/yolo_vision/tray_utils.py: classify_object, choose_slot',
      'tools/place_rounds.py: ClearTableSession.observe, known_inside, used_positions'])
s = new('Keep the motion sequence; update the destination XY', '04 / Reuse grasp & placement')
labels=['Approach','Descend\nto grasp','Lift','Transfer to\npre-release','Descend\nand release','Retreat','Return\nhome']
heights=[3.0,4.22,3.0,3.0,4.22,3.0,3.0]
for i in range(6):
    line(s,1.28+i*1.79,heights[i],1.28+(i+1)*1.79,heights[i+1],TEAL if i>1 else BLUE,2.3,True)
for i,(label,yy) in enumerate(zip(labels,heights)):
    xx=1.28+i*1.79
    disk(s,xx,yy,.24,BLUE if i<3 else TEAL,str(i+1),WHITE)
    text(s,label,xx-.72,4.84,1.44,.77,15,INK,True,PP_ALIGN.CENTER)
text(s,'Same rim contacts and finger direction. Same seven motion segments.',.76,1.87,11.9,.36,20,BLUE,True)
line(s,.91,4.59,12.38,4.59,LINE,1.3)
shape(s,.74,6.04,11.88,.68,PALE)
text(s,'tool release XY = selected slot XY + fixed offset',.99,6.22,7.18,.35,20,INK,True)
text(s,'Same Z • orientation • descent • home',8.46,6.25,3.9,.28,14,TEAL,True)
note(s,'Keep the motion sequence; update the destination XY',
     '原动作顺序继续使用：先到碗上方，再下降夹住、抬起、搬到放置准备点、下降释放、上退，最后回到保存的起点。一共七段，实际路径由原来的规划系统计算。图里的折线只表示动作先后，并不是机械臂真的走这样一条折线。\n薄托盘可以沿用原桌面释放高度，所以本次主要更换水平位置，也就是左右和前后的位置。要注意，夹爪的位置和碗最后落下的中心不一定重合，因此可以通过一次实际测量设置固定的水平偏移。这里没有针对每个抓取候选重新计算复杂的三维放置姿态。原 blue 基准的高度、方向、下降距离和返回姿态在任务开始时固定，整轮任务都沿用它们；实际落点误差仍要接设备后测量。',
     ['docs/圆碗径向抓取修改_20260928.md','tools/place_rounds.py: _release_coordinates, tray_release_config','src/grasp_executor/grasp_executor/grasp.py','src/grasp_executor/config/place_zones.json'])

s = new('Confirm a changed scene before starting the next pick', '05 / Placement verification')
tabletop(s,.73,2.17,5.29,2.54,1,1,1)
tabletop(s,7.16,2.17,5.29,2.54,2,0)
text(s,'BEFORE THE ACTION',.83,1.82,5.0,.28,12,GRAY,True)
text(s,'FRESH OBSERVATION AFTER RETURN',7.23,1.82,5.1,.28,12,TEAL,True)
line(s,6.21,3.4,6.92,3.4,TEAL,2.4,True)
checks=[('1','Source bowl is gone'),('2','New occupancy at the chosen slot'),('3','Previous tray occupancy remains')]
for i,(num,label) in enumerate(checks):
    x=.83+i*4.15
    tag(s,num,x,5.14,.4,TEAL,PALE)
    text(s,label,x,5.69,3.74,.85,20,INK,True)
text(s,'An action acknowledgment reserves the slot. A fresh observation confirms placement.',
     .82,6.64,11.9,.29,15,GRAY)
note(s,'Confirm a changed scene before starting the next pick',
     '机器人报告“动作执行完了”，只表示控制流程走完，并不自动证明碗落到了预想的位置。所以程序把两件事分开记：第一件是动作完成，先把那个位置保留；第二件是下一张图确认归位，才增加真正确认收好的数量。\n核查时看三点：原来的源位置是否还留着那只碗、选中的空位附近是否出现了新的碗、托盘里原有的碗是否还能对应上。三点一起判断，比只看到托盘里“有一只碗”可靠，因为那只碗可能本来就在里面。只要上一只的结果不清楚，就有限次重新观察，不继续搬下一只。图示为了容易理解只画了一只待收碗，实际程序每轮可以观察多只，但只执行一只。',
     ['tools/place_rounds.py: ClearTableSession.observe, record_motion','tools/test_clear_table_policy.py'])

s = new('Three placements, five observations', '06 / Recorded software result',
        'Plots reconstructed from saved synthetic observations • all sensing and actuator feedback are simulated')
text(s,'Blue: pending bowls   •   Teal: inside tray   •   Amber ring: selected destination',
     .76,1.64,11.9,.26,13,GRAY)
steps=['Pick → B','Confirm 1\nPick → C','Confirm 2\nPick → D','Confirm 3\nCheck empty','Confirm empty\nComplete']
for i,label in enumerate(steps):
    x=.62+i*2.55
    shape(s,x,2.01,2.38,3.77,WHITE,LINE)
    text(s,f'OBSERVATION {i+1}',x+.14,2.19,2.1,.25,11,TEAL,True)
    photo(s,f'observation_{i+1}.png',x+.08,2.54,2.22,2.18)
    text(s,label,x+.14,4.97,2.1,.64,16,INK,True,PP_ALIGN.CENTER)
shape(s,.65,6.1,12.0,.62,PALE)
text(s,'3 bowls confirmed   /   Slots B, C, D used   /   Existing slot A preserved',.91,6.27,11.48,.32,20,TEAL,True)
note(s,'The offline run completed three placements in five observations',
     '这一页不是随意画的流程，而是从本次保存的离线演示结果重新绘制的。蓝色圆表示托盘外的待收碗，绿色圆表示托盘内的碗，黄色虚线圈标出当轮选中的位置。所有图都使用同一坐标范围，所以能看出托盘在轮次间有轻微移动和转动。\n第一张图选 B 放第一只；第二张图先确认第一只，再选 C；第三张图确认第二只，再选 D。第四张图确认第三只已经进托盘，并第一次看到待收区域为空；第五张图再次确认空了，任务才结束。最初 A 位置已有一只碗，因此一直没有被重复使用。结果是三次抓放、三只确认归位、五次观察。这里输入图像位置和机械臂反馈都是模拟的，不能当作真实抓取成功率。',
     ['debug_output/clear_table_validation/yolo_20261010/offline_demo/result.json',
      'debug_output/clear_table_validation/yolo_20261010/offline_demo/round_0001..0005/synthetic_observation.json'])

s = new('Four checks protect detection and placement decisions', '07 / Debugging & regression checks')
cards=[('01','Model contract','A bowl model has no trained tray class.',
        'Check local weights, segmentation output and tray class.'),
       ('02','Complete view','A partial outline can suggest a misleading free area.',
        'Check depth, measured size, corners and the full view.'),
       ('03','Stable slot direction','A rectangular mask can flip its reported axis by 180°.',
        'Keep the closest direction; test occupied-slot continuity.'),
       ('04','Bowl clearance','A large bowl needs more space than a small one.',
        'Use the actual footprint; retain previously occupied slots.')]
for i,(num,title,problem,fix) in enumerate(cards):
    x=.73+(i%2)*6.31;y=1.97+(i//2)*2.38
    shape(s,x,y,5.96,2.1,WHITE,LINE);shape(s,x,y,.065,2.1,TEAL)
    text(s,num,x+.2,y+.18,.4,.3,13,TEAL,True)
    text(s,title,x+.8,y+.15,4.9,.38,22,INK,True)
    text(s,problem,x+.22,y+.75,5.48,.48,16,GRAY)
    text(s,fix,x+.22,y+1.38,5.48,.6,16,TEAL,True)
note(s,'Four checks protect detection and placement decisions', '',
     ['debug_output/clear_table_validation/yolo_20261010/summary.json','tools/test_clear_table_runner.py','tools/test_clear_table_policy.py','src/yolo_vision/test/test_tray_utils.py'])


s = new('Software checks passed; hardware trials are next', '08 / Validation & next step',
        'Recorded validation: 10 OCT 2026 • no physical camera or robot connected')
shape(s,.74,1.98,3.43,3.3,NAVY)
text(s,str(SUMMARY['pytest']['total_passed']),1.03,2.25,2.83,1.02,66,WHITE,True)
text(s,'automated tests passed',1.03,3.44,2.89,.7,21,WHITE,True)
text(s,f"{SUMMARY['pytest']['unit_regression_passed']} regression + 7 ROS tests",1.03,4.39,2.9,.4,15,'B6CFDA')
text(s,'3 packages built',1.03,4.88,2.9,.3,16,'60CEC3',True)
entries=[('14','ROS transport','Fresh messages and task identities'),
         ('32','Launch contracts','Topics, namespaces and parameters'),
         ('2','Exit cleanup','No owned workers left running'),
         ('16','MoveIt integration','Planning and simulated motion')]
for i,(count,title,body) in enumerate(entries):
    x=4.64+(i%2)*4.03; y=2.02+(i//2)*1.65
    text(s,count,x,y,.84,.65,35,TEAL,True)
    text(s,title,x+.96,y+.08,2.85,.38,19,INK,True)
    text(s,body,x,y+.82,3.65,.63,16,GRAY)
shape(s,.75,5.58,11.88,1.18,PALE)
text(s,'Before real execution',.97,5.79,3.45,.4,20,TEAL,True)
text(s,'Tray model → calibrate → preview → test one bowl',4.54,5.8,7.8,.39,17,INK,True)
text(s,'Thin, flat tray • stationary during motion • no new 3D collision scene',.97,6.37,11.37,.29,15,GRAY)
note(s,'Software checks passed; hardware trials are next', '',
     ['debug_output/clear_table_validation/yolo_20261010/summary.json',
      'debug_output/clear_table_validation/yolo_20261010/validation.log',
      'New_run/README.md', 'src/yolo_vision/config/clear_table_tray.json'])

short_notes = [
    '目标是把桌上的碗逐只放进可以移动的薄托盘。机器人每次只做一轮：看清桌面、量出位置、确认上一只、选择下一只和空位、按原方法抓放、回到起点。下一轮重新拍照，因为托盘和占用情况可能已经变化。第一轮没有上一只需要核查；最后一只放完，也要观察确认，不能仅凭动作结束就宣布收拾完成。托盘可以在两轮之间移动，动作期间要保持不动。',
    '现在直接用 YOLO 识别托盘，不贴标记。分割模型像是在照片上描出托盘轮廓，再用同一帧深度和坐标转换，把像素换成实际的厘米位置。蓝色框是外轮廓，用实测外尺寸核对；绿色框是可放碗的内侧范围。公式表示托盘上的固定放置点随托盘转动和平移。原碗模型继续工作。特别要说明：现有模型没有托盘类别，需要另外提供训练好的托盘分割权重，当前图是原理示意，并非实际识别截图。',
    '左边解决“抓哪只”：整只碗的占地范围都在托盘内，才算已经收好；它只占位置，不再抓。托盘外、且处于允许收拾区域内的碗，才是候选。压着边界或看不清时先重新观察。右边解决“放哪里”：预先设置几个宽松位置，按当前这只碗的大小检查离边缘和其他碗是否够远。刚用过的位置会保留，不能因为下一张图漏检就当成空位。新任务再根据现场观察重建占用。',
    '碗怎么夹继续沿用原来的碗沿抓法，手指方向也不变。动作仍是七段：到碗上方、下降夹住、抬起、搬到放置准备点、下降释放、上退、回起点。图只表达先后顺序，不是真实机械臂轨迹。因为托盘很薄，沿用原桌面释放高度，主要更新左右、前后的位置。夹爪与碗中心的差别用固定水平偏移补偿；没有为每个抓取候选重算复杂放置姿态。原高度和姿态在任务开始时读取一次。',
    '“动作完成”和“确认放好”是两回事。动作完成后先保留选中的位置；返回起点拍新图后，再检查三点：源位置的碗消失了、选中的空位出现新碗、托盘里原有的碗仍能对应上。这样可以避免把托盘里本来就有的碗误当作刚放下的碗。确认成功才增加归位数量，并继续下一只。结果不明确就有限次重新观察，不直接继续搬。',
    '这五张图来自保存的模拟运行数据，使用相同坐标范围。蓝色是待收碗，绿色是托盘内的碗，黄色圈是选中的位置。第一轮选 B，第二轮确认第一只并选 C，第三轮确认第二只并选 D。第四轮确认第三只并首次看到源区域为空，第五轮再次确认空了才结束。原来已有碗的 A 位一直没有重复使用。托盘在轮次间也有轻微移动。结果为三只归位、五次观察；这里的感知和机械臂反馈均为模拟。',
    '这次重点检查四件事。第一，文件存在还不够，模型必须真的有托盘类别和轮廓输出；缺少时直接报错。第二，不能把只看见的一部分当成整个托盘，要核对完整边界、尺寸、深度和视野。第三，矩形长轴可能被程序反过来表示，已补上相邻轮次的方向衔接，避免空位编号突然对调。第四，放置间隔继续按当前碗的实际大小检查。测试使用合成输入；实际大幅转动或替换托盘无法靠矩形轮廓可靠区分。',
    f"三个 ROS 软件包构建成功，{SUMMARY['pytest']['unit_regression_passed']} 项回归与 7 项独立 ROS 控制测试通过，合计 {SUMMARY['pytest']['total_passed']} 项。右边另列通信、启动、退出清理和 MoveIt 规划检查，不代表实机成功率。现在没有连接设备，也未提供训练好的托盘模型；下一步先准备模型、量好尺寸和偏移，做规划预览，再从一只碗开始实测。本版要求单个长短边可区分的矩形薄托盘，两轮间小幅移动，动作期间不动；未新增三维碰撞场景。"
]
for idx,(slide,body) in enumerate(zip(prs.slides,short_notes)):
    title, previous = notes[idx]
    sources = previous.split('【资料来源】\n',1)[1]
    value = '【中文讲解备注｜面向非专业听众】\n' + body + '\n\n【资料来源】\n' + sources
    slide.notes_slide.notes_text_frame.text = value
    for paragraph in slide.notes_slide.notes_text_frame.paragraphs:
        paragraph.font.name = 'Noto Sans CJK SC'; paragraph.font.size = Pt(12)
    notes[idx] = (title,value)

assert len(prs.slides) == 8
prs.core_properties.title = 'YOLO Movable-Tray Table Clearing — 8-Slide Methods Brief'
path = OUT / 'Movable_Tray_Table_Clearing_8slides_EN_CN.pptx'
prs.save(path)
with (OUT/'中文讲解备注_8页.txt').open('w',encoding='utf-8') as f:
    f.write('YOLO可移动薄托盘收桌面：8页英文PPT中文讲解\n建议时长：约8–10分钟。\n讲解备注已写入PPT各页备注区；托盘定位图已更新为YOLO分割原理示意。\n')
    for idx,(title,body) in enumerate(notes,1):
        f.write(f'\n第{idx:02d}页  {title}\n{body}\n')
(OUT/'asset_sources.json').write_text(json.dumps({
    'source_deck': 'output/movable_tray_presentation_20261010/Movable_Tray_Table_Clearing_EN_CN.pptx',
    'source_slide_mapping': [[1,3],[2,5,6],[3,7,8],[4,9,10],[5,11],[6,12],[7,13],[8,14,15]],
    'diagrams': 'Editable native PowerPoint shapes; schematic, not physical trajectories',
    'tray_localization': 'YOLO instance segmentation + RGB-D; editable schematic, no trained tray weights provided',
    'plots': 'Reconstructed from the current YOLO-version offline synthetic observations',
    'evidence': str(EVIDENCE.relative_to(ROOT)),
    'hardware_validated': False
}, ensure_ascii=False, indent=2),encoding='utf-8')
print(path)
